from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


TERMS = [
    "atributos",
    "pontos de vida",
    "pontos de mana",
    "pericias",
    "perícias",
    "tipos de ações",
    "ação padrão",
    "ação de movimento",
    "ação completa",
    "ação livre",
    "reação",
    "magias",
    "habilidades",
    "condições",
    "carga",
    "inventário",
    "defesa",
    "deslocamento",
]


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def analyze_xlsx(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        non_empty = []
        formulas = 0
        strings = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text = str(value)
                non_empty.append(cell.coordinate)
                if text.startswith("="):
                    formulas += 1
                elif len(text.strip()) > 0:
                    strings.append({"cell": cell.coordinate, "value": normalize(text)[:90]})

        sheets.append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "non_empty_count": len(non_empty),
            "formula_count": formulas,
            "merged_ranges": [str(rng) for rng in list(ws.merged_cells.ranges)[:20]],
            "sample_labels": strings[:80],
        })
    return {"file": str(path), "sheets": sheets}


def page_hits(reader: PdfReader) -> list[dict]:
    results = []
    for idx, page in enumerate(reader.pages):
        text = normalize(page.extract_text() or "")
        folded = text.lower()
        hits = []
        for term in TERMS:
            pos = folded.find(term)
            if pos == -1:
                continue
            start = max(0, pos - 80)
            end = min(len(text), pos + 220)
            hits.append({
                "term": term,
                "context": text[start:end],
            })
        if hits:
            results.append({"page": idx + 1, "hits": hits[:5]})
    return results


def analyze_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    first_pages = []
    for idx, page in enumerate(reader.pages[:4]):
        first_pages.append({
            "page": idx + 1,
            "sample": normalize(page.extract_text() or "")[:900],
        })

    return {
        "file": str(path),
        "pages": len(reader.pages),
        "first_pages": first_pages,
        "term_hits": page_hits(reader)[:60],
    }


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    if len(sys.argv) != 4:
        raise SystemExit("usage: analyze_tormenta_sources.py <xlsx> <book_pdf> <sheet_pdf>")

    xlsx_path = Path(sys.argv[1])
    book_pdf = Path(sys.argv[2])
    sheet_pdf = Path(sys.argv[3])

    data = {
        "xlsx": analyze_xlsx(xlsx_path),
        "book_pdf": analyze_pdf(book_pdf),
        "sheet_pdf": analyze_pdf(sheet_pdf),
    }
    print(json.dumps(data, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
